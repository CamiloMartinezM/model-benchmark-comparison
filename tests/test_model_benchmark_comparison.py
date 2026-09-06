"""Test cost calculations, comparability, and editable workbook behavior."""

from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
from pathlib import Path
from types import ModuleType

import numpy as np
import pandas as pd
import pytest
from matplotlib.figure import Figure
from openpyxl import load_workbook as read_excel_file

SCRIPT = Path(__file__).resolve().parents[1] / "model_benchmark_comparison.py"


@pytest.fixture
def app() -> ModuleType:
    """Load the script without invoking its command-line entry point.

    Returns:
        ModuleType: The imported benchmark comparison module.
    """
    assert SCRIPT.is_file(), "The standalone analysis script has not been implemented."
    spec = importlib.util.spec_from_file_location("model_benchmark_comparison", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_pareto_handles_ties_and_nonmonotonic_effort(app: ModuleType) -> None:
    """Reject costlier equal scores and retain equivalent efficient points.

    Args:
        app: The imported analysis module.
    """
    costs = np.array([4.0, 2.0, 3.0, 2.0, 1.0, np.nan])
    scores = np.array([60.0, 70.0, 70.0, 70.0, 50.0, 99.0])
    assert app.pareto_mask(costs, scores).tolist() == [
        False,
        True,
        False,
        True,
        True,
        False,
    ]


def test_success_cost_uses_probability_and_preserves_missing(app: ModuleType) -> None:
    """Catch percentage scaling errors and fabricated missing cost values.

    Args:
        app: The imported analysis module.
    """
    frame = pd.DataFrame(
        {
            "score": [50.0, 0.0, 80.0, 50.0],
            "cost_usd": [2.0, 2.0, np.nan, 2.0],
            "score_unit": ["percent", "percent", "percent", "index"],
            "cost_basis": ["usd_per_task"] * 4,
            "success_metric": [True, True, True, False],
            "comparison_group": ["same"] * 4,
        }
    )
    result = app.derive_metrics(frame)
    assert result.loc[0, "usd_per_success"] == 4.0
    assert pd.isna(result.loc[1, "usd_per_success"])
    assert pd.isna(result.loc[2, "usd_per_success"])
    assert pd.isna(result.loc[3, "usd_per_success"])


def test_pareto_does_not_mix_evaluation_groups(app: ModuleType) -> None:
    """Keep evaluations with distinct harnesses out of each other's frontier.

    Args:
        app: The imported analysis module.
    """
    frame = pd.DataFrame(
        {
            "score": [90.0, 50.0],
            "cost_usd": [1.0, 20.0],
            "score_unit": ["percent", "percent"],
            "cost_basis": ["usd_per_task", "usd_per_task"],
            "success_metric": [True, True],
            "comparison_group": ["harness_a", "harness_b"],
        }
    )
    assert app.derive_metrics(frame)["pareto_efficient"].tolist() == [True, True]


def test_workbook_roundtrip_retains_edits(app: ModuleType, tmp_path: Path) -> None:
    """Catch a rerun that replaces editable results with the embedded snapshot.

    Args:
        app: The imported analysis module.
        tmp_path: An isolated directory for the workbook.
    """
    path = tmp_path / "comparison.xlsx"
    app.ensure_workbook(path)
    workbook = read_excel_file(path)
    sheet = workbook["Results"]
    headings = [cell.value for cell in sheet[1]]
    score_column = headings.index("score") + 1
    sheet.cell(row=2, column=score_column).value = 42.125
    for name in ("score_lower", "score_upper"):
        sheet.cell(row=2, column=headings.index(name) + 1).value = None
    workbook.save(path)
    workbook.close()
    app.ensure_workbook(path)
    tables = app.load_workbook(path)
    assert tables["Results"].loc[0, "score"] == 42.125


def test_validator_rejects_invalid_score(app: ModuleType, tmp_path: Path) -> None:
    """Catch malformed Excel scores before they enter a plot.

    Args:
        app: The imported analysis module.
        tmp_path: An isolated directory for the workbook.
    """
    path = tmp_path / "comparison.xlsx"
    app.ensure_workbook(path)
    tables = app.load_workbook(path)
    tables["Results"].loc[0, "score"] = 101.0
    with pytest.raises(ValueError, match="score"):
        app.validate_tables(tables)


def test_provisional_cost_cannot_dominate_measured_cost(app: ModuleType) -> None:
    """Keep ineligible estimates from displacing measured efficient results.

    Args:
        app: The imported analysis module.
    """
    frame = pd.DataFrame(
        {
            "score": [90.0, 50.0],
            "cost_usd": [1.0, 20.0],
            "score_unit": ["percent", "percent"],
            "cost_basis": ["usd_per_task", "usd_per_task"],
            "success_metric": [True, True],
            "comparison_group": ["same", "same"],
            "efficiency_eligible": [False, True],
        }
    )
    assert app.derive_metrics(frame)["pareto_efficient"].tolist() == [False, True]


def test_boolean_text_preserves_false(app: ModuleType) -> None:
    """Catch text flags that Python would otherwise treat as true.

    Args:
        app: The imported analysis module.
    """
    assert app.parse_boolean("FALSE") is False
    assert app.parse_boolean("True") is True
    with pytest.raises(ValueError, match="boolean"):
        app.parse_boolean("unknown")


def test_script_runs_without_research_files(tmp_path: Path) -> None:
    """Generate a workbook and figures from a copy of the script alone.

    Args:
        tmp_path: An isolated directory without research source files.
    """
    copied = tmp_path / "standalone.py"
    shutil.copy2(SCRIPT, copied)
    output = tmp_path / "output"
    result = subprocess.run(
        [
            sys.executable,
            str(copied),
            "--output-dir",
            str(output),
            "--no-html",
            "--dpi",
            "50",
        ],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        check=False,
        timeout=60,
    )
    assert result.returncode == 0, result.stderr
    for name in (
        "model_benchmarks.xlsx",
        "model_benchmark_comparison.png",
        "analysis.csv",
    ):
        assert (output / name).stat().st_size > 1000


def uncertainty_rows() -> pd.DataFrame:
    """Build one renderable observation with a wide confidence interval.

    Returns:
        pd.DataFrame: A point at 50 with bounds at 40 and 60.
    """
    return pd.DataFrame(
        [
            {
                "model": "GPT-6 Astra",
                "effort": "high",
                "panel": "terminal",
                "score": 50.0,
                "score_lower": 40.0,
                "score_upper": 60.0,
                "cost_usd": 2.0,
                "cost_status": "owner_reported",
                "comparison_group": "test",
                "harness": "Test evaluation",
                "uncertainty": "Published 95% interval",
            }
        ]
    )


def test_filtered_figure_preserves_full_interval(app: ModuleType) -> None:
    """Catch axis limits that clip published uncertainty in a filtered view.

    Args:
        app: The imported analysis module.
    """
    axes = Figure().add_subplot()
    app.draw_panel(axes, uncertainty_rows(), app.PANELS[1], ("GPT-6 Astra",))
    lower, upper = axes.get_ylim()
    assert lower <= 40.0
    assert upper >= 60.0


def test_interactive_figure_contains_uncertainty(
    app: ModuleType, tmp_path: Path
) -> None:
    """Catch an HTML export that loses the static figure's uncertainty data.

    Args:
        app: The imported analysis module.
        tmp_path: An isolated output directory.
    """
    path = app.plot_interactive(uncertainty_rows(), tmp_path, ("GPT-6 Astra",))
    payload = path.read_text().split("const traces=", maxsplit=1)[1]
    traces, _ = json.JSONDecoder().raw_decode(payload)
    assert traces[0]["error_y"]["array"] == [10.0]
    assert traces[0]["error_y"]["arrayminus"] == [10.0]
    assert "Published 95% interval" in traces[0]["text"][0]
